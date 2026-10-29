"""在参考工作簿的副本上扩展机构/人员，验证所有其他 ZIP 成员不变。"""
from pathlib import Path
import json
import re
import sys
import zipfile

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
sys.path.insert(0, str(ROOT / 'tests'))
from group_msg_monitor.relay_core import (ENTRY, extract_people, read_roster, read_target, sheet_path,
                        update_file, file_hash, state_path)
from test_workflow import FUGUO, HUITIANFU, message


def main():
    original = ROOT / 'data/reference/要分名单 26Q3.xlsx'
    roster_path = ROOT / 'data/reference/研究所内部通讯录9.10.xlsx'
    output_dir = ROOT / '.work/reference-validation'
    output_dir.mkdir(parents=True, exist_ok=True)
    target_path = output_dir / '扩展样例.xlsx'
    # 允许重复执行验证，所有修改仅在独立副本中。
    state_path(target_path).unlink(missing_ok=True)
    original_hash = file_hash(original)
    roster = read_roster(roster_path)
    target = read_target(original)
    sample_names = sorted({n for e in ENTRY.findall(HUITIANFU + '\n' + FUGUO) for n in extract_people(e, roster)})
    columns = target['columns'].copy()
    next_col = max(columns.values()) + 1
    extra_headers = []
    for name in sample_names:
        if name not in columns:
            columns[name] = next_col
            extra_headers.append(f'<c r="{openpyxl.utils.get_column_letter(next_col)}1" t="inlineStr"><is><t>{name}</t></is></c>')
            next_col += 1
    with zipfile.ZipFile(original) as archive:
        path = sheet_path(archive, 'Sheet1')
        raw = archive.read(path).decode('utf-8')
        if extra_headers:
            first = re.search(r'<row\b[^>]*r="1"[^>]*>.*?</row>', raw, re.S)
            raw = raw[:first.start()] + first.group(0).replace('</row>', ''.join(extra_headers) + '</row>') + raw[first.end():]
        new_rows = []
        for row_no, inst in [(10, '富国基金（10.30）'), (11, '汇添富基金（11.6）')]:
            cells = [f'<c r="A{row_no}" t="inlineStr"><is><t>{inst}</t></is></c>']
            for name in sample_names:
                coord = f'{openpyxl.utils.get_column_letter(columns[name])}{row_no}'
                cells.append(f'<c r="{coord}" t="inlineStr"><is><t>扩展测试客户甲，客户乙</t></is></c>')
            new_rows.append(f'<row r="{row_no}">' + ''.join(cells) + '</row>')
        raw = raw.replace('</sheetData>', ''.join(new_rows) + '</sheetData>')
        raw = re.sub(r'<dimension ref="[^"]+"\s*/>', f'<dimension ref="A1:{openpyxl.utils.get_column_letter(next_col-1)}11"/>', raw)
        with zipfile.ZipFile(target_path, 'w') as dest:
            for item in archive.infolist():
                dest.writestr(item, raw.encode() if item.filename == path else archive.read(item.filename))
    result = update_file(target_path, roster_path, lambda *args: [message(HUITIANFU), message(FUGUO)],
                         first_start='2026-10-26', end='2026-10-26 09:00:00')
    assert result['cleared_count'] == 21, result
    assert result['issues'] == [], result['issues']
    with zipfile.ZipFile(result['backup']) as before, zipfile.ZipFile(target_path) as after:
        for member in before.namelist():
            if member != path:
                assert before.read(member) == after.read(member), member
    assert file_hash(original) == original_hash
    summary = {'roster_people': len(roster), 'original_target_people': len(target['columns']),
               'extra_people_for_test': sorted(set(sample_names) - set(target['columns'])),
               'cleared_count': result['cleared_count'], 'unchanged_original_sha256': original_hash,
               'result': result['archive'], 'non_target_zip_members': 'byte-for-byte unchanged'}
    (output_dir / 'verification.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
