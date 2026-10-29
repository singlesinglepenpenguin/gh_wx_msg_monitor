"""导出最近一次真实更新的核对表；不更新业务文件，不初始化微信。"""
from pathlib import Path
import sys
import json

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
import openpyxl
from openpyxl.styles import Alignment, Font
from group_msg_monitor.relay_core import summarize_issues, parse_bjt, latest_version, load_state
from group_msg_monitor.wechat_source import data_home

prefs = json.loads((data_home()/'preferences.json').read_text())
target = latest_version(prefs['target'])
folder = target.parent
state = load_state(target)
report = Path(state['last_archive']).parent/'details.json'
result = json.loads(report.read_text())
book = openpyxl.Workbook()
sheet = book.active
sheet.title = '本次结果'
sheet.append(['项目', '结果'])
for label, value in [('时间', result['updated_at']), ('群', result['group']),
                     ('新接龙', result['message_count']), ('历史复核', result['retried_pending_count']),
                     ('清空公式/值单元格', result['cleared_count']), ('原本已空', result['already_empty_count']),
                     ('待核对项（已合并重复）', len(summarize_issues(result['issues'])))]:
    sheet.append([label, value])
sheet = book.create_sheet('清空明细')
sheet.append(['机构', '姓名', '单元格', '接龙证据次数'])
for item in result['changes']:
    sheet.append([item['institution'], item['person'], item['cell'], len(item['evidence'])])
sheet = book.create_sheet('待核对')
sheet.append(['机构', '原因', '姓名或名单行', '出现次数'])
for item in summarize_issues(result['issues']):
    sheet.append([item['institution'], item['reason'], item['person'] or item['detail'], item['count']])
for sheet in book:
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width = 32
        for cell in col:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
path = folder/('核对报告_'+parse_bjt(result['updated_at']).strftime('%Y%m%d_%H%M%S')+'.xlsx')
book.save(path)
print(path)
