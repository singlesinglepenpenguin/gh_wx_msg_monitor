from __future__ import annotations

import hashlib
import json
import os
import re
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from group_msg_monitor.relay_core import (BJT, extract_people, file_hash, load_state, parse_bjt, latest_version,
                        plan_changes, read_roster, read_target, sheet_path, update_file, result_notes, write_run_note, latest_run_note, state_path, workbook_lock)


HUITIANFU = '''#接龙
🧧🧧🧧【汇添富权益】10月打分 今日10/26（周一）开启，截止11/6（下周五）
打派点参考服务记录，如需补充请联系我！！
1. 龚泓月 丁韵倩
2. 张雨蒙-金工 已要
3. 张婉姝 已要
4. 张建民@国海通信 已要
5. 匡培钦 交运 已要
6. 王鹏@国海化工&新材料 已要
7. 邱迪 已要
8. 胡家东@国海食饮 已要
9. 董伯骏-国海化工&新材料 ok
10. 刘熹@国海计算机 已要
11. 罗芳@国海 已要'''

FUGUO = '''#接龙
❗️【富国基金】9-10月开始集中打分🔥
富国基金权益#今天 开始集中打分，固收和量化部门周三开始集中打分，均于#本周五（10月30日）打完
建议大家今天盘后集中要分，不要分会漏打派点❗️
1. 山雨菲 已要
2. 杜先康 已要
3. 林加力 国海证券银行研究 已要
4. 邱迪 已要
5. 赵阳@国海策略 已要
6. 王鹏@国海化工&新材料 已要
7. 张建民@国海通信 已要
8. 董伯骏-国海化工&新材料 ok
9. 刘熹@国海计算机 已要'''


def message(content, when='2026-10-26 08:00:00'):
    return {'content': content, 'create_ts': int(parse_bjt(when).timestamp()), 'time': when,
            'sender': '测试发送者', 'chat_name': 'GH派点反馈群'}


def make_roster(path, names):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(['行业大类', '', '组别', '', '姓名'])
    for n in names:
        sheet.append(['研究', '', '测试组', '', n])
    book.save(path)


def make_target(path, names):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Sheet1'
    sheet.append(['机构/截止日期'] + list(names))
    for inst in ['汇添富基金（11.6）', '富国基金（10.30）', '鹏华基金（10.26）']:
        sheet.append([inst] + ['客户甲，客户乙'] * len(names))
    sheet['B2'] = '=1+2'
    sheet['B2'].number_format = '0.00'
    other = book.create_sheet('机构明细')
    other.append(['公式必须保留', '=SUM(1,2)'])
    sheet.freeze_panes = 'B2'
    book.save(path)


class WorkflowTests(unittest.TestCase):
    def test_chinese_excel_names_use_english_metadata_and_lock_names(self):
        target = self.root/'中文名单.xlsx'
        shutil.copyfile(self.target, target)
        with workbook_lock(target):
            locks = list(self.root.glob('.relay-lock-*.lock'))
            self.assertEqual(len(locks), 1)
            self.assertTrue(locks[0].name.isascii())
        result = update_file(target, self.roster, lambda *args: [message(HUITIANFU)], first_start='2026-10-26', end='2026-10-26 09:00:00')
        self.assertTrue(state_path(target).name.isascii())
        self.assertEqual(load_state(target)['output_name'], target.name)
        self.assertTrue((Path(result['archive']).parent/'details.json').is_file())
        for path in self.root.rglob('*'):
            if path.is_dir() or path.suffix != '.xlsx':
                self.assertTrue(path.name.isascii(), path)

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.roster = self.root / '通讯录.xlsx'
        self.target = self.root / '目标.xlsx'
        self.names = ['龚泓月', '丁韵倩', '张雨蒙', '张婉姝', '张建民', '匡培钦', '王鹏', '邱迪', '胡家东', '董伯骏', '刘熹', '罗芳', '山雨菲', '杜先康', '林加力', '赵阳']
        make_roster(self.roster, self.names)
        make_target(self.target, self.names)

    def run_update(self, messages, **kwargs):
        return update_file(self.target, self.roster, lambda *args: messages, first_start='2026-10-26 00:00:00', end='2026-10-26 09:00:00', **kwargs)

    def test_each_update_rereads_added_and_removed_institutions(self):
        self.run_update([])
        book = openpyxl.load_workbook(self.target)
        sheet = book['Sheet1']
        # Replace the prior institution after the first run; the next run
        # must use the changed workbook rather than its old institution list.
        sheet['A2'] = '新增测试基金'
        column = next(c.column for c in sheet[1] if c.value == '赵阳')
        sheet.cell(2, column, '待清空')
        book.save(self.target)
        book.close()
        result = update_file(self.target, self.roster, lambda *args: [
            message('#接龙\n【新增测试基金】\n1. 赵阳 已要', '2026-10-26 10:00:00'),
            message('#接龙\n【汇添富权益】\n1. 赵阳 已要', '2026-10-26 10:01:00'),
        ], end='2026-10-26 15:00:00')
        self.assertEqual(result['cleared_count'], 1)
        self.assertEqual(result['changes'][0]['institution'], '新增测试基金')
        self.assertEqual(result['issues'][0]['reason'], '机构未匹配')

    def test_real_appmsg_prefix_and_confirmed_institution_alias(self):
        book = openpyxl.load_workbook(self.target)
        book['Sheet1']['A2'] = '中邮基金（10.29）'
        book.save(self.target)
        text = '[appmsg type=53] #接龙\n【中邮创业基金】季度打分\n1. 王鹏@国海化工&新材料 已要'
        result = self.run_update([message(text)])
        self.assertEqual(result['cleared_count'], 1)
        self.assertEqual(result['changes'][0]['person'], '王鹏')
        self.assertEqual(result['issues'], [])

    def test_full_names_with_emoji_and_html_separators(self):
        roster = {name: ['测试组'] for name in ('钟琪', '赵博轩', '朱谞晟', '刘旭德', '赵兰亭')}
        for entry, person in [('钟琪🌟国海公用&amp;核聚变 已要', '钟琪'),
                              ('赵博轩✈️@国海军工 已要', '赵博轩'),
                              ('朱谞晟&amp;国海计算机 已要', '朱谞晟'),
                              ('刘旭德🍶国海食饮 已要', '刘旭德'),
                              ('赵兰亭🌟国海 已要', '赵兰亭')]:
            self.assertEqual(extract_people(entry, roster), [person])
        for entry in ('某钟琪🌟国海', '钟琪某 已要', '钟🌟琪 已要'):
            self.assertEqual(extract_people(entry, roster), [])

    def test_q3_heading_and_instruction_entry_are_handled(self):
        book = openpyxl.load_workbook(self.target)
        book['Sheet1']['A2'] = '信达澳亚（10.29）'
        book.save(self.target)
        text = '[appmsg type=53] #接龙\n信达澳亚Q3打分开启\n1. 龚泓月 已要\n2. 盘后请大家积极要分⛽️'
        result = self.run_update([message(text)])
        self.assertEqual(result['cleared_count'], 1)
        self.assertEqual(result['issues'], [])
        self.assertEqual(len(result['ignored_entries']), 1)

    def test_generated_versions_preserve_input_and_continue_without_stacked_timestamps(self):
        before = file_hash(self.target)
        first = self.run_update([message(HUITIANFU)], save_as_new=True)
        output = Path(first['output'])
        self.assertEqual(file_hash(self.target), before)
        self.assertFalse(state_path(self.target).exists())
        self.assertIn('_20261026_090000.xlsx', output.name)
        output_before = file_hash(output)
        second = update_file(output, self.roster, lambda *args: [], end='2026-10-26 09:05:00', save_as_new=True)
        self.assertEqual(file_hash(output), output_before)
        self.assertTrue(Path(second['output']).name.endswith('_20261026_090500.xlsx'))
        self.assertNotIn('090000', Path(second['output']).name)
        self.assertEqual(load_state(Path(second['output']))['through'], parse_bjt('2026-10-26 09:05:00').isoformat())
        second_hash = file_hash(Path(second['output']))
        third = update_file(second['output'], self.roster, lambda *args: [], end='2026-10-26 09:05:00', save_as_new=True)
        self.assertNotEqual(second['output'], third['output'])
        self.assertTrue(Path(third['output']).name.endswith('_20261026_090500_1.xlsx'))
        self.assertEqual(Path(third['notes_file']).name, 'update_20261026_090500_1.txt')
        self.assertEqual(file_hash(Path(second['output'])), second_hash)
        self.assertEqual(latest_version(self.target), Path(third['output']))
        book = openpyxl.load_workbook(self.target)
        book['Sheet1']['A2'] = '汇添富基金（新的运营批次）'
        book.save(self.target)
        self.assertEqual(latest_version(self.target), self.target.resolve())

    def test_fixed_output_has_one_timestamped_note_per_update_and_opens_latest(self):
        outputs, records = self.root/'outputs', self.root/'records'
        before = file_hash(self.target)
        first = self.run_update([message(HUITIANFU)], save_as_new=True, output_dir=outputs, records_dir=records)
        self.assertEqual(Path(first['output']).parent, outputs.resolve())
        self.assertEqual(file_hash(self.target), before)
        self.assertEqual(latest_version(self.target, outputs), Path(first['output']))
        second = update_file(first['output'], self.roster, lambda *args: [message(HUITIANFU, '2026-10-26 09:01:00')],
                             end='2026-10-26 09:05:00', save_as_new=True, output_dir=outputs, records_dir=records)
        self.assertEqual(second['cleared_count'], 0)
        self.assertEqual(len(second['already_empty_entries']), second['already_empty_count'])
        self.assertGreater(second['already_empty_count'], 0)
        for cell in second['already_empty_entries']:
            self.assertIn('单元格：'+cell['cell'], result_notes(second))
            self.assertIn('姓名：'+cell['person'], result_notes(second))
        self.assertIn('本轮没有删除任何单元格', result_notes(second))
        first_note, second_note = Path(first['notes_file']), Path(second['notes_file'])
        self.assertEqual(first_note.name, 'update_20261026_090000.txt')
        self.assertEqual(second_note.name, 'update_20261026_090500.txt')
        self.assertNotEqual(first_note, second_note)
        self.assertEqual(latest_run_note(records), second_note)
        self.assertEqual(second_note.read_text(encoding='utf-8').count('更新时间（北京时间）'), 1)
        self.assertIn(first['output'], first_note.read_text(encoding='utf-8'))
        self.assertIn('已删除内容', first_note.read_text(encoding='utf-8'))
        self.assertIn('原本已空', second_note.read_text(encoding='utf-8'))
        self.assertFalse((records/'all_updates.txt').exists())
        self.assertEqual(len(list(outputs.glob('*.xlsx'))), 2)

    def test_repeated_relay_evidence_counts_as_one_cell(self):
        entries = [dict(message('#接龙\n【富国基金】\n1. 赵阳 已要', f'2026-10-26 08:0{i}:00'), sender=sender)
                   for i, sender in enumerate(('发送者甲', '发送者乙', '发送者丙'))]
        result = self.run_update(entries)
        self.assertEqual(result['cleared_count'], 1)
        self.assertEqual(len(result['changes'][0]['evidence']), 3)
        self.assertIn('3 条接龙包含该姓名', result_notes(result))
        self.assertIn('本轮只删除 1 格', result_notes(result))

    def test_failed_version_write_does_not_leave_empty_generated_file(self):
        before = file_hash(self.target)
        with patch('group_msg_monitor.relay_core.patch_workbook', side_effect=OSError('fixture write failure')):
            with self.assertRaisesRegex(OSError, 'fixture write failure'):
                self.run_update([message(HUITIANFU)], save_as_new=True)
        self.assertEqual(file_hash(self.target), before)
        self.assertFalse(list(self.root.glob(self.target.stem+'_20261026_*.xlsx')))

    def test_plain_institution_heading_before_instruction_brackets(self):
        book = openpyxl.load_workbook(self.target)
        book['Sheet1']['A2'] = '国联安（10.30）'
        book.save(self.target)
        text = '[appmsg type=53] #Group Note\n#接龙\n‼️国联安9-10月服务打分开启\n【本周五】前完成\n1. 龚泓月 已要'
        result = self.run_update([message(text)])
        self.assertEqual(result['cleared_count'], 1)
        self.assertEqual(result['issues'], [])

    def test_missing_institutions_report_heading_without_body_brackets(self):
        result = self.run_update([
            message('#接龙\n【建信基金】9-10月打分开始\n维持排名[抱拳]\n1. 龚泓月 已要'),
            message('#Group Note\n#接龙\n国联安9-10月服务打分开启\n【本周五】前完成\n1. 龚泓月 已要'),
        ])
        self.assertEqual(result['cleared_count'], 0)
        self.assertEqual([i['detail'] for i in result['issues']], ['建信基金', '国联安'])
        self.assertEqual(len(load_state(self.target)['pending']), 2)
        note = result_notes(result)
        self.assertIn('[INSTITUTION_MATCH]', note)
        self.assertIn('本轮实际读取的机构列：', note)

    def test_body_cannot_redirect_an_unknown_title_to_an_existing_institution(self):
        result = self.run_update([message('#接龙\n【未知基金】打分开始\n可参考【富国基金】服务记录\n1. 龚泓月 已要')])
        self.assertEqual(result['cleared_count'], 0)
        self.assertEqual(result['issues'][0]['detail'], '未知基金')

    def test_broker_mention_in_body_is_not_a_fund_institution(self):
        book = openpyxl.load_workbook(self.target)
        book['Sheet1']['A2'] = '国海基金（10.28）'
        book.save(self.target)
        text = '#接龙\n【其他基金】季度打分\n请将【国海证券】提名前三\n1. 龚泓月 已要'
        result = self.run_update([message(text)])
        self.assertEqual(result['cleared_count'], 0)
        self.assertEqual(result['issues'][0]['reason'], '机构未匹配')

    def test_styled_empty_cells_do_not_consume_following_cells(self):
        with zipfile.ZipFile(self.target) as archive:
            sheet = sheet_path(archive, 'Sheet1')
            raw = archive.read(sheet)
            for coord in (b'B2', b'C2'):
                raw = re.sub(rb'<c\b[^>]*\br="'+coord+rb'"[^>]*>.*?</c>',
                             b'<c r="'+coord+b'" s="1"/>', raw, count=1, flags=re.S)
            members = [(item, raw if item.filename == sheet else archive.read(item.filename)) for item in archive.infolist()]
        with zipfile.ZipFile(self.target, 'w') as dest:
            for item, data in members:
                dest.writestr(item, data)
        before = file_hash(self.target)
        result = self.run_update([message('#接龙\n【汇添富基金】\n1. 龚泓月')])
        self.assertEqual(result['cleared_count'], 0)
        self.assertEqual(result['already_empty_count'], 1)
        self.assertEqual(file_hash(self.target), before)

    def test_user_examples_and_cumulative_lists(self):
        original = self.target.read_bytes()
        # 最新版故意删除罗芳，较早版本出现过仍算完成。
        newer = HUITIANFU.rsplit('\n11.', 1)[0]
        result = self.run_update([message(HUITIANFU), message(newer, '2026-10-26 08:10:00'), message(FUGUO)])
        self.assertEqual(result['cleared_count'], 21)
        self.assertEqual(result['relay_count'], 3)
        self.assertEqual(result['issues'], [])
        self.assertEqual(Path(result['backup']).read_bytes(), original)
        self.assertEqual(Path(result['archive']).read_bytes(), self.target.read_bytes())
        book = openpyxl.load_workbook(self.target)
        self.assertIsNone(book['Sheet1'].cell(2, 13).value)  # 罗芳
        self.assertEqual(book['Sheet1']['B4'].value, '客户甲，客户乙')
        self.assertEqual(book['Sheet1']['B2'].number_format, '0.00')
        self.assertEqual(book['Sheet1'].freeze_panes, 'B2')
        book.close()
        with zipfile.ZipFile(result['backup']) as before, zipfile.ZipFile(self.target) as after:
            changed_sheet = sheet_path(before, 'Sheet1')
            self.assertEqual(before.namelist(), after.namelist())
            for name in before.namelist():
                if name != changed_sheet:
                    self.assertEqual(before.read(name), after.read(name), name)

    def test_repeat_and_next_window_are_idempotent(self):
        first = self.run_update([message(FUGUO)])
        second = update_file(self.target, self.roster, lambda *args: [message(FUGUO, '2026-10-26 09:00:00')],
                             end='2026-10-26 15:00:00')
        self.assertEqual(first['cleared_count'], 9)
        self.assertEqual(second['cleared_count'], 0)
        self.assertEqual(second['already_empty_count'], 9)
        self.assertEqual(load_state(self.target)['through'], '2026-10-26T15:00:00+08:00')
        self.assertNotEqual(first['archive'], second['archive'])

    def test_fetch_failure_leaves_original_and_progress(self):
        fingerprint = file_hash(self.target)
        def fail(*args):
            raise ValueError('读取失败')
        with self.assertRaisesRegex(ValueError, '读取失败'):
            update_file(self.target, self.roster, fail, first_start='2026-10-26', end='2026-10-26 09:00:00')
        self.assertEqual(file_hash(self.target), fingerprint)
        self.assertEqual(load_state(self.target), {})
        self.assertFalse(self.target.with_name('.' + self.target.name + '.relay.lock').exists())

    def test_state_failure_does_not_skip_on_retry(self):
        from group_msg_monitor.relay_core import atomic_json
        def failing_state(path, value):
            if Path(path).name.startswith('.relay-state-'):
                raise OSError('state write failure')
            return atomic_json(path, value)
        with patch('group_msg_monitor.relay_core.atomic_json', side_effect=failing_state):
            with self.assertRaisesRegex(OSError, 'state write failure'):
                self.run_update([message(FUGUO)])
        self.assertEqual(load_state(self.target), {})
        second = self.run_update([message(FUGUO)])
        self.assertEqual(second['cleared_count'], 0)
        self.assertEqual(second['already_empty_count'], 9)

    def test_no_ordinary_chat_and_no_partial_name(self):
        roster = read_roster(self.roster)
        self.assertEqual(extract_people('赵阳阳 已要', roster), [])
        self.assertEqual(extract_people('某赵阳 已要', roster), [])
        result = self.run_update([message('普通聊天 赵阳 富国基金 已要'), message('#接龙\n【富国基金】\n1. 赵阳阳 已要')])
        self.assertEqual(result['cleared_count'], 0)
        self.assertEqual(result['issues'], [])
        self.assertEqual(len(result['ignored_entries']), 1)

    def test_multiple_names_and_no_status_requirement(self):
        roster = read_roster(self.roster)
        self.assertEqual(extract_people('龚泓月 丁韵倩', roster), ['龚泓月', '丁韵倩'])
        result = self.run_update([message('#接龙\n【富国基金】\n1. 赵阳\n2. 王鹏@国海化工&新材料')])
        self.assertEqual(result['cleared_count'], 2)

    def test_reorder_and_add_rows_columns(self):
        book = openpyxl.load_workbook(self.target)
        sheet = book['Sheet1']
        sheet.insert_cols(2)
        sheet['B1'] = '新增人员'
        sheet.insert_rows(2)
        sheet['A2'] = '新增机构（10.31）'
        sheet['B2'] = '保留'
        book.save(self.target)
        book.close()
        result = self.run_update([message(FUGUO)])
        self.assertEqual(result['cleared_count'], 9)
        self.assertTrue(all(x['cell'].endswith('4') for x in result['changes']))

    def test_unknown_institution_reported_but_missing_person_skipped(self):
        result = self.run_update([message('#接龙\n【未知基金】\n1. 赵阳'), message('#接龙\n【富国基金】\n1. 通讯录没有此人')])
        self.assertEqual(result['cleared_count'], 0)
        self.assertEqual(len(result['issues']), 1)
        self.assertEqual(len(result['ignored_entries']), 1)

    def test_file_changed_during_read_not_overwritten(self):
        replacement = b'changed by another program'
        def fetch(*args):
            self.target.write_bytes(replacement)
            return [message(FUGUO)]
        with self.assertRaisesRegex(ValueError, '其他程序修改'):
            update_file(self.target, self.roster, fetch, first_start='2026-10-26', end='2026-10-26 09:00:00')
        self.assertEqual(self.target.read_bytes(), replacement)
        self.assertEqual(load_state(self.target), {})

    def test_future_messages_do_not_affect_earlier_cutoff(self):
        result = self.run_update([message(FUGUO, '2026-10-26 15:01:00')])
        self.assertEqual(result['cleared_count'], 0)

    def test_duplicate_headers_stop_update(self):
        book = openpyxl.load_workbook(self.target)
        book['Sheet1']['B1'] = book['Sheet1']['C1'].value
        book.save(self.target)
        book.close()
        with self.assertRaisesRegex(ValueError, '重复姓名'):
            self.run_update([message(FUGUO)])

    def test_missing_people_are_not_retried_after_roster_changes(self):
        first = self.run_update([message('#接龙\n【富国基金】\n1. 新同事')])
        self.assertEqual(first['issues'], [])
        self.assertEqual(len(first['ignored_entries']), 1)
        self.assertEqual(load_state(self.target)['pending'], [])
        make_roster(self.roster, self.names + ['新同事'])
        book = openpyxl.load_workbook(self.target)
        column = book['Sheet1'].max_column + 1
        book['Sheet1'].cell(1, column, '新同事')
        book['Sheet1'].cell(3, column, '待清空')
        book.save(self.target)
        book.close()
        second = update_file(self.target, self.roster, lambda *args: [], end='2026-10-26 15:00:00')
        self.assertEqual(second['cleared_count'], 0)
        self.assertEqual(load_state(self.target)['pending'], [])

    def test_missing_target_column_is_ignored_while_valid_people_are_processed(self):
        make_roster(self.roster, self.names + ['新同事'])
        result = self.run_update([message('#接龙\n【富国基金】\n1. 新同事 已要\n2. 赵阳 已要')])
        self.assertEqual(result['cleared_count'], 1)
        self.assertEqual(result['issues'], [])
        self.assertEqual(result['ignored_entries'][0]['reason'], '目标表没有此人列')
        self.assertEqual(load_state(self.target)['pending'], [])


if __name__ == '__main__':
    unittest.main()
