"""接龙匹配与 Excel 更新；仅修改目标 worksheet，保留其他 ZIP 成员。"""
from __future__ import annotations

import hashlib
import html
import json
import os
import re
import shutil
import tempfile
import unicodedata
import zipfile
from collections import defaultdict
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path
from xml.etree import ElementTree as ET
import posixpath

import openpyxl

BJT = timezone(timedelta(hours=8))
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
TAG = "{" + NS["m"] + "}"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def now_bjt():
    return datetime.now(BJT).replace(microsecond=0)


def parse_bjt(value):
    if isinstance(value, datetime):
        dt = value
    else:
        dt = datetime.fromisoformat(str(value).strip())
    return dt.replace(tzinfo=BJT) if dt.tzinfo is None else dt.astimezone(BJT)


def atomic_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp = tempfile.mkstemp(dir=path.parent, prefix=".writing-", suffix=".json")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as stream:
            json.dump(value, stream, ensure_ascii=False, indent=2)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temp, path)
    finally:
        if os.path.exists(temp):
            os.unlink(temp)


def file_hash(path):
    h = hashlib.sha256()
    with open(path, "rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def norm(value):
    return unicodedata.normalize("NFKC", str(value or "")).strip()


def read_roster(path):
    """按列名定位，不读取或保存电话、邮箱；组别继承合并单元格。"""
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    people = defaultdict(set)
    try:
        for sheet in book:
            name_col = group_col = header_row = None
            for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 30)), 1):
                labels = {norm(c.value): c.column for c in row if c.value is not None}
                if "姓名" in labels:
                    name_col, group_col, header_row = labels["姓名"], labels.get("组别"), row_number
                    break
            if name_col is None:
                continue
            group = ""
            for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
                if group_col and len(row) >= group_col and row[group_col - 1]:
                    group = norm(row[group_col - 1])
                if len(row) >= name_col and row[name_col - 1]:
                    person = re.sub(r"\s+", "", norm(row[name_col - 1]))
                    if re.fullmatch(r"[\u3400-\u9fff·]{2,12}", person):
                        people[person].add(group)
    finally:
        book.close()
    if not people:
        raise ValueError("通讯录未找到「姓名」列或有效姓名，请检查所选文件。")
    return {name: sorted(groups) for name, groups in people.items()}


def institution_key(value):
    value = norm(value)
    value = re.sub(r"[（(][^()（）]*[)）]", "", value)
    value = re.sub(r"\s+", "", value)
    value = re.sub(r"(?:权益|固收|量化)(?:部门|部)?$", "", value)
    value = re.sub(r"基金(?:管理)?(?:有限责任公司|有限公司)?$", "", value)
    # 已确认的同机构简称；不做子串或近似机构猜测。
    return {"中邮创业": "中邮"}.get(value, value)


def sheet_path(archive, name):
    book = ET.fromstring(archive.read("xl/workbook.xml"))
    rels = {r.attrib["Id"]: r.attrib["Target"] for r in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))}
    for sheet in book.findall("m:sheets/m:sheet", NS):
        if sheet.attrib["name"] == name:
            target = rels[sheet.attrib["{" + REL + "}id"]]
            return target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
    raise ValueError(f"目标文件没有工作表「{name}」。")


def read_target(path, sheet_name="Sheet1"):
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet_name not in book:
            raise ValueError(f"目标文件没有工作表「{sheet_name}」。")
        sheet = book[sheet_name]
        header_row = inst_col = None
        for row in sheet.iter_rows(min_row=1, max_row=min(20, sheet.max_row)):
            hits = [c for c in row if "机构" in norm(c.value) and ("姓名" in norm(c.value) or "截止" in norm(c.value))]
            if hits:
                header_row, inst_col = hits[0].row, hits[0].column
                break
        if header_row is None:
            raise ValueError("Sheet1 未找到「机构/截止日期」表头。")
        columns = defaultdict(list)
        for row in sheet.iter_rows(min_row=header_row, max_row=header_row):
            for cell in row:
                if cell.column != inst_col and cell.value:
                    columns[norm(cell.value)].append(cell.column)
        duplicate = [n for n, cs in columns.items() if len(cs) > 1]
        if duplicate:
            raise ValueError("目标表存在重复姓名列：" + "、".join(duplicate))
        rows = {}
        for row in sheet.iter_rows(min_row=header_row + 1):
            cell = row[inst_col - 1]
            if cell.value:
                label = norm(cell.value)
                key = institution_key(label)
                if key in rows:
                    raise ValueError(f"目标表机构重复或无法区分：{label}")
                rows[key] = {"row": cell.row, "label": label}
        if not rows or not columns:
            raise ValueError("目标表没有机构或姓名列。")
        return {"rows": rows, "columns": {n: cs[0] for n, cs in columns.items()}, "sheet": sheet_name}
    finally:
        book.close()


ENTRY = re.compile(r"(?m)^\s*\d+\s*[.、．)）]\s*(.+)$")
TITLE = re.compile(r"【([^】\n]+)】|\[([^\]\n]+)\]")
SKIPPED_PERSON_REASONS = {'通讯录未匹配姓名', '目标表没有此人列'}


def extract_people(entry, roster):
    """仅完整姓名匹配，允许多个姓名；不把短姓名匹配到较长姓名内部。"""
    names = sorted(roster, key=lambda x: (-len(x), x))
    if not names:
        return []
    text = norm(html.unescape(entry))
    # 表情作为分隔符；不能直接拼接两侧汉字，以免把“钟🌟琪”猜成“钟琪”。
    text = re.sub(r'[\U0001F000-\U0001FAFF\u2300-\u23FF\u2600-\u27BF\uFE0E\uFE0F\u200D\u20E3]', ' ', text)
    pattern = r"(?:^|[\s，,、;；:：/&])(" + "|".join(map(re.escape, names)) + r")(?=$|[\s，,、;；:：/&@＠\-—–(（]|已要|已完成|已联系|收到|[Oo][Kk])"
    return list(dict.fromkeys(m.group(1) for m in re.finditer(pattern, text)))


def heading_institutions(head, target):
    """没有机构括号时，仅匹配标题首行的完整机构名，不扫描正文。"""
    for line in head.splitlines():
        line = re.sub(r'^\[appmsg type=\d+\]\s*', '', line).strip()
        if not line or line.startswith('#'):
            continue
        line = re.sub(r'^[^\u3400-\u9fffA-Za-z]+', '', line)
        candidates = set()
        for key, row in target['rows'].items():
            name = re.sub(r'[（(][^()（）]*[)）]', '', norm(row['label'])).strip()
            if line.startswith(name):
                tail = line[len(name):]
                if not tail or re.match(r'[\s\d（(：:，,。!！—\-]|[Qq][1-4]|服务|打分|季度|开始|截止', tail):
                    candidates.add(key)
        return candidates
    return set()


def plan_changes(messages, roster, target):
    changes, issues, pending, ignored = {}, [], [], []
    relay_count = 0
    for message in sorted(messages, key=lambda m: m.get("create_ts", 0)):
        text = norm(html.unescape(message.get("content", "")))
        text = re.sub(r'^\[appmsg type=\d+\]\s*', '', text)
        if "#接龙" not in text:
            continue
        relay_count += 1
        issue_count = len(issues)
        # 每条消息独立定位机构，不从上一条消息继承。
        head = text[:ENTRY.search(text).start()] if ENTRY.search(text) else text
        titles = [a or b for a, b in TITLE.findall(head)]
        candidates = {institution_key(t) for t in titles if institution_key(t) in target["rows"]}
        if not candidates:
            candidates = heading_institutions(head, target)
        evidence = {k: message.get(k, "") for k in ("time", "create_ts", "sender", "chat_name")}
        if len(candidates) != 1:
            issues.append({**evidence, "reason": "机构未匹配" if not candidates else "一条接龙匹配多个机构", "detail": "、".join(titles), "content": text})
            pending.append(message)
            continue
        inst = next(iter(candidates))
        context = {**evidence, 'institution': target['rows'][inst]['label']}
        entries = ENTRY.findall(text)
        if not entries:
            issues.append({**evidence, "reason": "没有编号名单", "content": text})
        for entry in entries:
            names = extract_people(entry, roster)
            if not names:
                if re.match(r'^(?:盘后请(?:大家|老师)|最后一天请各位领导|主观打分[,，]|很多客户已提前打分)', entry):
                    ignored.append({**context, 'entry': entry, 'reason': '提示语，不参与姓名匹配'})
                    continue
                ignored.append({**context, 'entry': entry, 'reason': '通讯录未匹配姓名'})
                continue
            for name in names:
                if len(roster[name]) > 1:
                    issues.append({**context, "reason": "通讯录同名人员需核对", "detail": entry, "person": name})
                    continue
                if name not in target["columns"]:
                    ignored.append({**context, 'entry': entry, 'person': name, 'reason': '目标表没有此人列'})
                    continue
                coord = f"{openpyxl.utils.get_column_letter(target['columns'][name])}{target['rows'][inst]['row']}"
                changes.setdefault(coord, {"cell": coord, "person": name, "institution": target["rows"][inst]["label"], "evidence": []})["evidence"].append({**evidence, "entry": entry})
        if len(issues) > issue_count:
            pending.append(message)
    return {"relay_count": relay_count, "changes": list(changes.values()), "issues": issues, "pending": pending, "ignored": ignored}


def patch_workbook(source, output, plan, sheet_name="Sheet1"):
    """只删命中格的公式/值，其他工作表和缓存逐字节保留。"""
    changes = {c["cell"]: c for c in plan["changes"]}
    cleared, already_empty = [], []
    with zipfile.ZipFile(source) as archive:
        path = sheet_path(archive, sheet_name)
        raw = archive.read(path)
        # 保留原 namespace 前缀、extLst、格式与未命中格；避免 ElementTree 重写整表。
        seen = set()
        cell_pattern = re.compile(rb'<(?:\w+:)?c\b[^>]*?(?:/>|>.*?</(?:\w+:)?c\s*>)', re.S)
        def replace(match):
            cell = match.group(0)
            attr = re.search(rb'\br="([A-Z]+\d+)"', cell)
            if not attr or attr.group(1).decode() not in changes:
                return cell
            coord = attr.group(1).decode()
            seen.add(coord)
            root = ET.fromstring(b'<wrapper xmlns="' + NS['m'].encode() + b'">' + cell + b'</wrapper>')[0]
            children = [c for c in root if c.tag.rsplit('}', 1)[-1] in ('f', 'v', 'is')]
            old = [{"kind": c.tag.rsplit('}', 1)[-1], "text": ''.join(c.itertext())} for c in children]
            if children:
                cleared.append({**changes[coord], "previous": old})
            else:
                already_empty.append(coord)
            cleaned = re.sub(rb'<(?:\w+:)?(?:f|v|is)\b[^>]*?(?:/>|>.*?</(?:\w+:)?(?:f|v|is)\s*>)', b'', cell, flags=re.S)
            # 空格不再携带字符串类型，原样式保留。
            return re.sub(rb'\s+t="[^"]*"', b'', cleaned, count=1)
        modified = cell_pattern.sub(replace, raw)
        already_empty += sorted(set(changes) - seen)
        with zipfile.ZipFile(output, 'w') as dest:
            for item in archive.infolist():
                dest.writestr(item, modified if item.filename == path else archive.read(item.filename))
    with zipfile.ZipFile(output) as archive:
        if archive.testzip():
            raise ValueError("生成文件 ZIP 校验失败。")
        ET.fromstring(archive.read(path))
    return cleared, already_empty


@contextmanager
def workbook_lock(source):
    lock = source.with_name('.relay-lock-' + workbook_id(source) + '.lock')
    try:
        fd = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
    except FileExistsError:
        raise ValueError("此文件正在更新。若上次意外退出，请确认工具已关闭后处理同目录 .relay-lock-*.lock 文件。")
    try:
        os.write(fd, str(os.getpid()).encode())
        os.close(fd)
        yield
    finally:
        lock.unlink(missing_ok=True)


def workbook_id(source):
    return hashlib.sha256(Path(source).name.encode('utf-8')).hexdigest()[:16]


def state_path(source):
    source = Path(source)
    return source.with_name('.relay-state-' + workbook_id(source) + '.json')


def load_state(source):
    p = state_path(Path(source))
    if not p.exists():
        # 兼容迁移前的进度；新写入始终使用英文哈希文件名。
        p = Path(source).with_name('.' + Path(source).name + '.relay-state.json')
        if not p.exists():
            return {}
    try:
        return json.loads(p.read_text('utf-8'))
    except (OSError, ValueError) as exc:
        raise ValueError("更新进度文件损坏，请先恢复，避免跳过消息。") from exc


def latest_version(source, output_dir=None):
    """沿生成记录接续最新版本；输入文件已被人工改动时不跳转。"""
    source = Path(source).resolve()
    if not source.is_file():
        return source
    versions = []
    roots = {source.parent, Path(output_dir).resolve()} if output_dir else {source.parent}
    for path in (p for folder in roots for pattern in ('.relay-state-*.json', '.*.xlsx.relay-state.json') for p in folder.glob(pattern)):
        try:
            state = json.loads(path.read_text('utf-8'))
            name = state.get('output_name')
            if not name and path.name.endswith('.xlsx.relay-state.json'):
                name = path.name[1:-len('.relay-state.json')]
            if not name:
                continue
            output = path.with_name(name)
            if state.get('input_source') and state.get('through') and output.is_file():
                versions.append((output, state))
        except (OSError, ValueError):
            continue
    visited = set()
    while source not in visited:
        visited.add(source)
        children = [(path, state) for path, state in versions if state['input_source'] == str(source)]
        if not children:
            break
        fingerprint = file_hash(source)
        children = [(path, state) for path, state in children if state.get('input_sha256') == fingerprint]
        if not children:
            break
        source, _ = max(children, key=lambda item: (item[1]['through'], item[0].name))
    return source


def summarize_issues(issues):
    grouped = {}
    for issue in issues:
        if issue.get('reason') in SKIPPED_PERSON_REASONS:
            continue
        key = tuple(issue.get(k, '') for k in ('reason', 'institution', 'person', 'detail'))
        item = grouped.setdefault(key, {**dict(zip(('reason', 'institution', 'person', 'detail'), key)), 'count': 0})
        item['count'] += 1
    return list(grouped.values())


def result_notes(result):
    issues = summarize_issues(result['issues'])
    lines = [f"更新时间（北京时间）：{parse_bjt(result['through']):%Y-%m-%d %H:%M:%S}",
             f"群：{result['group']}",
             f"基于文件：{result['source']}",
             f"生成文件：{result.get('output', result['source'])}",
             f"上次成功进度：{result.get('previous_through') or '首次运行或旧记录未提供'}",
             f"本次读取 {result['message_count']} 条接龙；复核历史 {result['retried_pending_count']} 条。",
             f"清空 {result['cleared_count']} 格；原本已空 {result['already_empty_count']} 格。",
             '按机构、姓名和单元格合并，多条接龙证据不会重复计算删除格数。',
             '消息发送者是发布整条接龙的账号，不代表其中每个名单成员。',
             '缺少姓名、缺少人员列及纯提示语直接跳过，不列入待核对。', '', '本次实际删除：']
    def details(entries, action):
        for c in entries:
            lines.append(f"{action} | 机构：{c['institution']} | 姓名：{c['person']} | 单元格：{c['cell']}")
            if c.get('previous'):
                lines.append('  删除前：' + ('公式（完整公式及缓存值保存在 details.json）' if any(p['kind'] == 'f' for p in c['previous']) else '有内容的单元格（原值保存在 details.json）'))
            evidence = c.get('evidence', [])
            if len(evidence) > 1:
                lines.append(f"  {len(evidence)} 条接龙包含该姓名，合并为同一个单元格；" + ('本轮只删除 1 格。' if action == '已删除内容' else '本轮未再次删除。'))
            for e in evidence:
                lines.append(f"  接龙时间：{e.get('time', '')} | 消息发送者：{e.get('sender', '')} | 名单原文：{e.get('entry', '')}")
    details(result['changes'], '已删除内容')
    if not result['changes']:
        lines.append('无。本轮没有删除任何单元格内容。')
    lines += ['', '匹配成功但原本已空（本轮未再次删除）：']
    details(result.get('already_empty_entries', []), '原本已空')
    if not result.get('already_empty_entries'):
        lines.append('无。' if not result['already_empty_count'] else '旧版记录未保存此项人员明细，不能仅凭计数还原。')
    lines += ['', '正常跳过（无需处理）：']
    for entry in result.get('ignored_entries', []):
        lines.append(f"机构：{entry.get('institution', '')} | 姓名/原文：{entry.get('person') or entry.get('entry', '')} | 原因：{entry.get('reason', '')} | 时间：{entry.get('time', '')}")
    if not result.get('ignored_entries'):
        lines.append('无。')
    visible_records = sum(i.get('reason') not in SKIPPED_PERSON_REASONS for i in result['issues'])
    lines += ['', f"待核对 {len(issues)} 项（来自 {visible_records} 条记录，已合并重复；完整证据保存在details.json）："]
    lines += [f"{i['institution']} {i['reason']}：{i['person'] or i['detail']}（出现 {i['count']} 次）" for i in issues]
    if issues:
        lines += ['', '待核对原始证据（未据此删除单元格）：']
        for issue in result['issues']:
            if issue.get('reason') in SKIPPED_PERSON_REASONS:
                continue
            lines.append(f"时间：{issue.get('time', '')} | 发送人：{issue.get('sender', '')} | 原因：{issue.get('reason', '')}")
            lines.append(issue.get('content') or issue.get('detail') or issue.get('person', ''))
    return '\n'.join(lines)


def write_run_note(result, records_dir):
    """一轮一份英文名 TXT；时间及同秒版本序号与 Excel 对应。"""
    records_dir = Path(records_dir).resolve()
    records_dir.mkdir(parents=True, exist_ok=True)
    output = result.get('output', result['source'])
    match = re.search(r'_(\d{8}_\d{6}(?:_\d+)?)$', Path(output).stem)
    stamp = match.group(1) if match else parse_bjt(result['through']).strftime('%Y%m%d_%H%M%S')
    index = 0
    while True:
        suffix = '' if index == 0 else f'_{index}'
        destination = records_dir / f'update_{stamp}{suffix}.txt'
        try:
            fd = os.open(destination, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
        except FileExistsError:
            if f'生成文件：{output}' in destination.read_text(encoding='utf-8').splitlines():
                break
            index += 1
        else:
            os.close(fd)
            break
    fd, temporary = tempfile.mkstemp(dir=records_dir, prefix='.records-', suffix='.txt')
    try:
        with os.fdopen(fd, 'w', encoding='utf-8') as f:
            f.write(result_notes(result))
        os.replace(temporary, destination)
    finally:
        Path(temporary).unlink(missing_ok=True)
    return destination


def latest_run_note(records_dir):
    runs = []
    for path in Path(records_dir).glob('*/details.json'):
        result = json.loads(path.read_text(encoding='utf-8'))
        output = Path(result.get('output', result['source']))
        if output.is_file() and load_state(output).get('last_archive') == result.get('archive'):
            runs.append((result['through'], path.stat().st_mtime_ns, result))
    if not runs:
        return None
    return write_run_note(max(runs, key=lambda item: (item[0], item[1]))[2], records_dir)


def reserve_version(source, stamp, family):
    """预留独立的新版本，避免同秒或重复运行覆盖已有文件。"""
    index = 0
    while True:
        suffix = '' if index == 0 else f'_{index}'
        path = source.with_name(f'{family}_{stamp}{suffix}.xlsx')
        try:
            fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
        except FileExistsError:
            index += 1
            continue
        os.close(fd)
        return path


def update_file(source, roster_path, fetch, *, first_start=None, end=None, group="GH派点反馈群", log=lambda s: None, save_as_new=False, output_dir=None, records_dir=None):
    source, roster_path = Path(source).resolve(), Path(roster_path).resolve()
    if not source.is_file() or source.suffix.lower() != '.xlsx':
        raise ValueError("请选择有效的目标 .xlsx 文件。")
    if source == roster_path:
        raise ValueError("通讯录和目标文件不能是同一个文件。")
    if source.with_name('~$' + source.name).exists():
        raise ValueError("Excel 正在使用此文件，请关闭后再更新。")
    with workbook_lock(source):
        state = load_state(source)
        end = parse_bjt(end or now_bjt())
        if state and state.get('group') != group:
            raise ValueError("当前进度属于另一个群，请使用新的目标文件。")
        start_value = state.get('through') or first_start
        if not start_value:
            raise ValueError("首次运行请填写开始时间（北京时间），之后会自动接续。")
        start = parse_bjt(start_value)
        if start > end:
            raise ValueError("开始时间晚于结束时间。")
        fingerprint = file_hash(source)
        roster = read_roster(roster_path)
        target = read_target(source)
        log(f"已读取通讯录 {len(roster)} 人，目标表 {len(target['rows'])} 家机构、{len(target['columns'])} 个人名列。")
        # 重读边界一分钟，避免同秒消息和上次读取期间的到达差异；清空操作幂等。
        query_start = start - timedelta(minutes=1) if state else start
        log(f"读取 {query_start:%Y-%m-%d %H:%M:%S} 至 {end:%Y-%m-%d %H:%M:%S}（北京时间）。")
        messages = fetch(query_start, end, group)
        if file_hash(source) != fingerprint:
            raise ValueError("读取期间目标 Excel 被其他程序修改，本次没有覆盖原文件，请重新运行。")
        messages = [m for m in messages if int(query_start.timestamp()) <= m.get('create_ts', 0) <= int(end.timestamp())]
        fresh_count = len(messages)
        combined = {}
        for msg in state.get('pending', []) + messages:
            key = (msg.get('create_ts', 0), msg.get('sender', ''), msg.get('content', ''))
            combined[key] = msg
        messages = list(combined.values())
        plan = plan_changes(messages, roster, target)
        stamp = end.strftime('%Y%m%d_%H%M%S')
        family = state.get('version_family') or re.sub(r'_\d{8}_\d{6}(?:_\d+)?$', '', source.stem)
        run_id = stamp + '_' + os.urandom(3).hex()
        history = (Path(records_dir).resolve() if records_dir else source.parent / 'update_records') / run_id
        history.mkdir(parents=True, exist_ok=False)
        backup = history / ((family + '_' + stamp + '_更新前.xlsx') if save_as_new else source.stem + '_更新前.xlsx')
        shutil.copy2(source, backup)
        fd, tmp = tempfile.mkstemp(dir=source.parent, prefix='.relay-', suffix='.xlsx')
        os.close(fd)
        output, published = source, False
        try:
            if save_as_new:
                destination = Path(output_dir).resolve() if output_dir else source.parent
                destination.mkdir(parents=True, exist_ok=True)
                output = reserve_version(destination / source.name, stamp, family)
            archive_path = history / (output.name if save_as_new else source.stem + '_' + stamp + '.xlsx')
            cleared, empty = patch_workbook(backup, tmp, plan, target['sheet'])
            empty_cells = set(empty)
            result_hash = file_hash(tmp)
            shutil.copy2(tmp, archive_path)
            result = {"status": "saved", "updated_at": end.isoformat(), "from": query_start.isoformat(), "through": end.isoformat(), "group": group,
                      "source": str(source), "archive": str(archive_path), "backup": str(backup), "message_count": fresh_count,
                      "output": str(output), "previous_through": state.get('through'),
                      "retried_pending_count": len(state.get('pending', [])),
                      "relay_count": plan['relay_count'], "cleared_count": len(cleared), "already_empty_count": len(empty),
                      "changes": cleared, "already_empty_entries": [c for c in plan['changes'] if c['cell'] in empty_cells],
                      "issues": plan['issues'], "result_sha256": result_hash,
                      "ignored_entries": plan['ignored']}
            # 在覆盖前完成归档和审计；状态失败时不推进，下次可以幂等重放。
            atomic_json(history / 'details.json', result)
            if file_hash(source) != fingerprint:
                raise ValueError("读取期间目标 Excel 被其他程序修改，本次没有覆盖原文件，请重新运行。")
            os.replace(tmp, output)
            published = True
            new_state = {"through": end.isoformat(), "group": group, "last_archive": str(archive_path), "result_sha256": result_hash, "pending": plan['pending'], "version_family": family, "output_name": output.name}
            if save_as_new:
                new_state.update(input_source=str(source), input_sha256=fingerprint)
            atomic_json(state_path(output), new_state)
            try:
                result['notes_file'] = str(write_run_note(result, history.parent))
            except OSError as exc:
                log('Excel 已保存，但本轮 TXT 未能保存，可在“查看更新记录”重试：' + str(exc))
            log(f"更新完成：清空 {len(cleared)} 格，未匹配 {len(plan['issues'])} 项。归档：{archive_path}")
            return result
        finally:
            Path(tmp).unlink(missing_ok=True)
            if save_as_new and output != source and not published:
                output.unlink(missing_ok=True)
